"""Build the exercise dataset for the 'Organising data for AI' session.

Synthetic Durban beach water-quality data. One 'truth' table is generated,
then rendered (a) as a tidy CSV and (b) as a deliberately messy workbook that
contains every common mistake from the Data Carpentry spreadsheet lesson.
"""
import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).parent  # files are written next to this script
OUT.mkdir(exist_ok=True)
rng = np.random.default_rng(1110)

GUIDE = 185  # illustrative threshold used only for the red-cell trap
DL, UPPER = 10, 2000  # detection limit and upper countable limit (cfu/100 mL)


def make_truth():
    rows = []
    plan = {
        2024: ("Sipho", ["North Beach", "Addington"], dt.date(2024, 3, 5)),
        2025: ("Lerato", ["North Beach", "Addington", "uShaka", "Umhlanga Main"],
               dt.date(2025, 3, 4)),
    }
    sid = 0
    for year, (who, sites, start) in plan.items():
        for site in sites:
            for wk in range(8):
                sid += 1
                d = start + dt.timedelta(days=7 * wk)
                val = float(np.round(rng.lognormal(np.log(70), 1.1)))
                below = above = False
                miss = ""
                if val < DL:
                    below, val = True, np.nan
                elif val > UPPER:
                    above, val = True, np.nan
                rows.append(dict(
                    sample_id=f"DBN-{year}-{sid:03d}", site=site,
                    sample_date=d, collector=who,
                    enterococci_cfu_100ml=val, below_detection=below,
                    above_count_range=above,
                    water_temp_c=round(float(rng.normal(22.5, 1.2)), 1),
                    missing_reason=miss, notes=""))
    df = pd.DataFrame(rows)
    # Plant specific cases so every trap is present and reproducible.
    def setrow(i, **kw):
        for k, v in kw.items():
            df.at[i, k] = v
    setrow(2, enterococci_cfu_100ml=np.nan, missing_reason="bottle_broken")
    setrow(5, enterococci_cfu_100ml=np.nan, below_detection=True)
    setrow(9, enterococci_cfu_100ml=np.nan, above_count_range=True)
    setrow(12, enterococci_cfu_100ml=np.nan, missing_reason="not_sampled_rough_sea")
    setrow(14, notes="sample may have leaked in cooler box")
    setrow(20, enterococci_cfu_100ml=0.0, below_detection=False)  # true zero? no -> see dictionary
    setrow(20, enterococci_cfu_100ml=np.nan, below_detection=True)
    setrow(33, notes="after heavy rain")
    return df


truth = make_truth()

# ---------------- tidy outputs ----------------
tidy = truth.copy()
tidy["sample_date"] = pd.to_datetime(tidy["sample_date"]).dt.strftime("%Y-%m-%d")
tidy.to_csv(OUT / "beach_water_quality_tidy.csv", index=False)

dictionary = pd.DataFrame([
    ("sample_id", "Unique sample identifier, DBN-<year>-<sequence>", "text", "", "Never reused"),
    ("site", "Beach sampling site", "category", "",
     "North Beach | Addington | uShaka | Umhlanga Main"),
    ("sample_date", "Date the water sample was collected", "date", "ISO 8601 YYYY-MM-DD", ""),
    ("collector", "Field assistant who collected the sample", "category", "", "First name only (synthetic)"),
    ("enterococci_cfu_100ml", "Enterococci count", "number", "cfu per 100 mL",
     "Blank when not measured, below detection, or above countable range (see flags)"),
    ("below_detection", "Count below method detection limit", "boolean", "",
     f"TRUE means < {DL} cfu/100 mL; value column left blank"),
    ("above_count_range", "Plate too numerous to count", "boolean", "",
     f"TRUE means > {UPPER} cfu/100 mL; value column left blank"),
    ("water_temp_c", "Surface water temperature", "number", "degrees Celsius", ""),
    ("missing_reason", "Why no measurement exists", "category", "",
     "blank | bottle_broken | not_sampled_rough_sea"),
    ("notes", "Free-text field note", "text", "", "Plain text, no line breaks"),
], columns=["column", "description", "type", "units_or_format", "allowed_values_or_notes"])
dictionary.to_csv(OUT / "beach_water_quality_data_dictionary.csv", index=False)

# ---------------- messy workbook ----------------
wb = Workbook()
RED = PatternFill("solid", fgColor="FF9999")
YEL = PatternFill("solid", fgColor="FFF2A8")
BOLD = Font(bold=True)


def messy_value(r):
    if r.missing_reason == "bottle_broken":
        return -999
    if r.missing_reason == "not_sampled_rough_sea":
        return None  # blank: was it zero, missing or forgotten?
    if r.below_detection:
        return "<10"
    if r.above_count_range:
        return "TNTC"
    v = int(r.enterococci_cfu_100ml)
    return f"{v} cfu" if r.Index % 7 == 3 else v


def messy_date(d, i):
    styles = [lambda d: d,                                   # real Excel date
              lambda d: d.strftime("%d %b"),                 # no year
              lambda d: d.strftime("%Y/%m/%d"),
              lambda d: d.strftime("%d-%m-%Y"),
              lambda d: d.strftime("%m/%d/%Y")]              # US order
    return styles[i % len(styles)](d)


# --- Tab 1: Sipho 2024: two tables side by side, title rows, legend ---
ws = wb.active
ws.title = "2024 Sipho"
ws["A1"] = "DURBAN BEACH SAMPLING 2024 - Sipho"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:I1")
blocks = {"North Beach": 1, "Addington": 6}
for site, c0 in blocks.items():
    ws.cell(3, c0, site.upper()).font = BOLD
    ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 3)
    ws.cell(3, c0).alignment = Alignment(horizontal="center")
    for j, h in enumerate(["Date", "Entero (cfu/100mL)", "Temp °C", "Comments"]):
        ws.cell(4, c0 + j, h).font = BOLD
    sub = truth[(truth.collector == "Sipho") & (truth.site == site)]
    for k, r in enumerate(sub.itertuples()):
        row = 5 + k
        ws.cell(row, c0, messy_date(r.sample_date, k))
        cell = ws.cell(row, c0 + 1, messy_value(r))
        if isinstance(cell.value, (int, float)) and cell.value > GUIDE:
            cell.fill = RED
        if cell.value == "TNTC":
            cell.fill = RED
        ws.cell(row, c0 + 2, f"{r.water_temp_c}" if k % 3 else r.water_temp_c)
        if r.notes:
            ws.cell(row, c0 + 3, r.notes)
    ws.column_dimensions[ws.cell(4, c0 + 1).column_letter].width = 20
ws["B6"].comment = Comment("Retest? Lab said contamination possible", "Sipho")
ws["A15"] = "KEY: red = exceeds guideline!!  -999 = bottle broke.  Blank = ?"
ws["A15"].font = Font(italic=True)

# --- Tab 2: Lerato 2025: one long table, inconsistent everything ---
ws2 = wb.create_sheet("2025_Lerato (final) v2")
hdr = ["Site", "Sample date", "Enterococci", "Water temp", "Site / tide", "Notes"]
for j, h in enumerate(hdr, 1):
    ws2.cell(1, j, h).font = BOLD
spell = {"North Beach": ["North Beach", "north beach", "North Bch", "North Beach "],
         "Addington": ["Addington", "Addington Beach", "ADDINGTON"],
         "uShaka": ["uShaka", "Ushaka", "U-Shaka"],
         "Umhlanga Main": ["Umhlanga Main", "Umhlanga", "Umhlanga Rocks"]}
tides = ["HW+2 / calm", "LW, choppy", "mid tide", "HW-1 / swell 2m"]
sub = truth[truth.collector == "Lerato"].reset_index(drop=True)
row = 2
for k, r in enumerate(sub.itertuples()):
    if k == 16:  # repeated header row part-way down
        for j, h in enumerate(hdr, 1):
            ws2.cell(row, j, h).font = BOLD
        row += 1
    ws2.cell(row, 1, spell[r.site][k % len(spell[r.site])])
    ws2.cell(row, 2, messy_date(r.sample_date, k + 1))
    v = messy_value(r)
    ws2.cell(row, 3, v)
    if isinstance(v, (int, float)) and v > GUIDE:
        ws2.cell(row, 3).fill = RED
    t = r.water_temp_c
    ws2.cell(row, 4, f"{t}C" if k % 4 == 0 else (round(t * 9 / 5 + 32, 1) if k == 11 else t))
    ws2.cell(row, 5, tides[k % 4])
    note = r.notes or ("1M swimmer, 2F swimmers in water" if k == 6 else "")
    if k == 9:
        note = "check this one—entered from WhatsApp “photo”"
    ws2.cell(row, 6, note)
    if k in (3, 4):
        for j in range(1, 7):
            ws2.cell(row, j).fill = YEL  # yellow = 'exclude' known only to Lerato
    row += 1
    if k == 7:
        ws2.cell(row, 1, "Mean North Beach").font = BOLD
        ws2.cell(row, 3, "=AVERAGE(C2:C9)")
        row += 1
ws2.column_dimensions["F"].width = 40

# --- Tab 3: summary with formulas that silently include -999 ---
ws3 = wb.create_sheet("Summary")
ws3["A1"] = "Summary for report"; ws3["A1"].font = BOLD
ws3["A3"], ws3["B3"] = "Mean Entero North Beach 2024", "=AVERAGE('2024 Sipho'!B5:B12)"
ws3["A4"], ws3["B4"] = "Mean Entero Addington 2024", "=AVERAGE('2024 Sipho'!G5:G12)"
ws3["A5"], ws3["B5"] = "Samples exceeding (count red cells)", "Sipho to fill in"
ws3["A7"] = "Data is SYNTHETIC teaching data - not real monitoring results."
ws3["A7"].font = Font(italic=True, color="AA0000")
ws3.column_dimensions["A"].width = 40
wb.save(OUT / "beach_water_quality_MESSY.xlsx")

# ---------------- numbers for the facilitator demo ----------------
nb = truth[(truth.collector == "Sipho") & (truth.site == "North Beach")]
messy_nums = []
for r in nb.itertuples():
    v = messy_value(r)
    if isinstance(v, (int, float)):
        messy_nums.append(v)
true_vals = nb.enterococci_cfu_100ml.dropna()
print("Messy cells North Beach 2024:", [messy_value(r) for r in nb.itertuples()])
print("Excel AVERAGE of numeric cells (incl -999):", round(np.mean(messy_nums), 1))
print("Mean of measured values only:", round(true_vals.mean(), 1), "n =", len(true_vals))
print("Rows below detection:", int(nb.below_detection.sum()),
      "above range:", int(nb.above_count_range.sum()))
print(truth.groupby(["collector", "site"]).size())
